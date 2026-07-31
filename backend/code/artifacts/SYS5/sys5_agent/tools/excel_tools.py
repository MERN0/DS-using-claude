"""
Deterministic, no-judgment IO tools over the real client `.xlsx` files.

These are the only tools that ever touch the client's real input directory
or the output path. They do no domain reasoning (no guessing which sheet is
"the requirements sheet", no signal/command matching heuristics) -- every
judgment call about *what the data means* is left to the agent that calls
these tools. This mirrors how `ls`/`grep`/file-write tools work in Claude
Code: cheap, deterministic, and dumb on purpose.

All tools return JSON strings so results are stable across model providers
and easy for the agent to read back.

## Sandboxing

Every filesystem-touching tool here is built by a factory
(`build_read_only_tools` / `build_write_tool`) bound to exactly one run's
`input_folder_path` / final output file, rather than defined as bare
module-level functions that would accept any path the model hands them.
This mirrors the `virtual_mode` sandbox `FilesystemBackend` already applies
to the agent's built-in filesystem tools (see `agent/build.py`): a run may
only ever read from inside its own input directory and write to its own
predetermined output file, regardless of what path a (possibly confused or
prompt-injected) model turn tries to pass in. Build a fresh set of tool
instances per run -- never share one factory's output across two different
callers' directories.

Every read-only tool below takes a bare `file_name` (never a path) -- the
input directory itself is fixed per run via the closure, so the model is
never asked to know, construct, or re-type it. This is deliberate, not just
a convenience: an absolute path is exactly the kind of value a model can get
wrong across turns/agents (dropping the directory, mis-joining it, or
mangling a Windows-style path when the run happens to execute on a
different OS than whatever machine the path string originally came from),
and `list_input_files` already hands back the one `file_name` value that
every other tool here needs, verbatim.
"""

from __future__ import annotations

import json
import os
import re
import uuid
from pathlib import Path
from typing import Any, Optional

from langchain_core.tools import tool
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from sys5_agent.config import settings

# ---------------------------------------------------------------------------
# Output workbook styling (cosmetic only -- never touches cell content)
# ---------------------------------------------------------------------------

_HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
_HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
_HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=True)
_BODY_ALIGNMENT = Alignment(horizontal="left", vertical="top", wrap_text=True)
_ZEBRA_FILL = PatternFill(start_color="F2F6FA", end_color="F2F6FA", fill_type="solid")
_THIN_SIDE = Side(style="thin", color="B7C3CC")
_BORDER = Border(left=_THIN_SIDE, right=_THIN_SIDE, top=_THIN_SIDE, bottom=_THIN_SIDE)

# Per-column width in characters. Narrow identifier/enum-like columns stay
# compact; free-text columns get more room (and rely on wrap_text + row
# height rather than becoming unreadably wide).
_NARROW_COLUMNS = {"Test Case ID", "Variant", "Check Type", "Mode of Execution", "Priority"}
_WIDE_COLUMNS = {
    "Test Case Description",
    "Test Precondition",
    "Test Input Data",
    "Test Steps",
    "Expected Result",
}
_MIN_WIDTH = 14
_NARROW_WIDTH = 16
_WIDE_WIDTH = 48
_DEFAULT_WIDTH = 28
_ROW_HEIGHT = 90


def _beautify_workbook(ws: Worksheet, columns: list[str], row_count: int) -> None:
    """Apply a readable default style to the freshly written output sheet.

    Purely cosmetic: header styling, column widths sized to field type,
    wrapped/top-aligned body cells, thin borders, zebra striping, a frozen
    header row, and an autofilter -- none of it changes any cell value.
    """
    for col_idx, name in enumerate(columns, start=1):
        letter = get_column_letter(col_idx)
        header_cell = ws.cell(row=1, column=col_idx)
        header_cell.fill = _HEADER_FILL
        header_cell.font = _HEADER_FONT
        header_cell.alignment = _HEADER_ALIGNMENT
        header_cell.border = _BORDER

        if name in _WIDE_COLUMNS:
            width = _WIDE_WIDTH
        elif name in _NARROW_COLUMNS:
            width = _NARROW_WIDTH
        else:
            width = _DEFAULT_WIDTH
        ws.column_dimensions[letter].width = max(width, _MIN_WIDTH)

    for row_idx in range(2, row_count + 2):
        ws.row_dimensions[row_idx].height = _ROW_HEIGHT
        is_zebra = (row_idx % 2) == 0
        for col_idx in range(1, len(columns) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.alignment = _BODY_ALIGNMENT
            cell.border = _BORDER
            if is_zebra:
                cell.fill = _ZEBRA_FILL

    ws.freeze_panes = "A2"
    last_col_letter = get_column_letter(len(columns))
    ws.auto_filter.ref = f"A1:{last_col_letter}{row_count + 1}"


def _dump(obj: Any) -> str:
    return json.dumps(obj, ensure_ascii=False, indent=2, default=str)


def _cell_str(value: Any) -> Any:
    if value is None:
        return None
    return value


# ---------------------------------------------------------------------------
# Sandboxed path resolution
# ---------------------------------------------------------------------------


_PATH_SEP_RE = re.compile(r"[\\/]")


def _candidate_for(root: Path, raw: str) -> Optional[Path]:
    candidate = Path(raw)
    candidate = candidate if candidate.is_absolute() else (root / candidate)
    candidate = candidate.resolve()
    if candidate != root and root not in candidate.parents:
        return None
    return candidate


def _resolve_within(root: Path, file_name: str) -> Optional[Path]:
    """Resolve `file_name` against `root`, refusing to leave `root`.

    The tools below only ever hand the calling agent a bare file name (see
    `list_input_files`) -- `root` is already fixed per run via closure, so
    the model never needs to know or reconstruct an absolute path, and
    never gets a chance to garble one (e.g. a Windows-style
    "C:\\Users\\...\\x.xlsx" typed on a POSIX host resolves to nothing --
    backslashes aren't separators there -- which used to surface as a
    confusing "no such file or directory C:\\Users\\..." error). If a
    model still hands back something path-like anyway (its own guess, or a
    stray value copied from elsewhere), this falls back to just the last
    path segment (split on both '/' and '\\') before giving up, so that
    degrades gracefully instead of failing outright.

    Symlinks are resolved too, so a symlink inside `root` pointing outside
    it is also rejected. Returns `None` (never raises) when nothing
    resolves inside `root`, so callers can turn that into a normal
    tool-error JSON instead of crashing the agent turn.
    """
    root = root.resolve()
    candidate = _candidate_for(root, file_name)
    if candidate is not None and candidate.exists():
        return candidate

    tail = _PATH_SEP_RE.split(str(file_name).strip())[-1]
    if tail and tail != str(file_name):
        fallback = _candidate_for(root, tail)
        if fallback is not None and fallback.exists():
            return fallback

    return candidate


def _access_denied(file_name: str, root: Path) -> str:
    return _dump(
        {
            "error": (
                f"Access denied: '{file_name}' resolves outside the allowed "
                f"directory for this run ({root}). Pass just the bare file "
                "name (as returned by list_input_files), not a path."
            )
        }
    )


# ---------------------------------------------------------------------------
# Read-only tools (bound to one run's input directory)
# ---------------------------------------------------------------------------


def build_read_only_tools(input_root: Path) -> list:
    """Build the read-only excel tools, sandboxed to `input_root`.

    Every path any of these tools resolves is checked against `input_root`
    (see `_resolve_within`) -- a request for a file outside it comes back
    as a tool-error JSON, never an actual read. Call this once per run with
    that run's own input directory; never reuse the returned tools across a
    different run's directory.
    """
    root = Path(input_root).resolve()

    @tool
    def list_input_files() -> str:
        """List every .xlsx file found directly inside this run's input directory.

        Use this first, before anything else, to see what's actually in the
        client's data directory -- the requirements file plus an unknown mix
        of supporting files (signal lists, command lists, compound command
        lists, application parameters, communication matrices, or other
        reference data). File names are NOT standardized across clients; do
        not assume any particular name means a particular thing without
        previewing it. There is no directory to pass in -- this run's input
        directory is fixed, and every `file_name` this returns is exactly
        what every other tool here expects (no directory prefix needed).

        Returns:
            JSON list of {"file_name": str, "size_bytes": int}.
        """
        files = sorted(p for p in root.iterdir() if p.is_file() and p.suffix.lower() in (".xlsx", ".xlsm"))
        return _dump([{"file_name": p.name, "size_bytes": p.stat().st_size} for p in files])

    @tool
    def list_workbook_sheets(file_name: str) -> str:
        """List every sheet in a workbook along with its used dimensions.

        Call this on a file before previewing/reading it, so you know what
        sheets exist and roughly how large each one is (a sheet with 1 row is
        almost certainly not the requirements sheet; a sheet with 500+ rows
        and only 2-3 columns is more likely a flat signal/command list than a
        requirements sheet).

        Args:
            file_name: Exact file name, as returned by list_input_files (no
                directory -- this run's input directory is fixed).

        Returns:
            JSON list of {"sheet_name": str, "max_row": int, "max_col": int}.
        """
        p = _resolve_within(root, file_name)
        if p is None:
            return _access_denied(file_name, root)
        if not p.is_file():
            return _dump({"error": f"Not a file: {p}"})
        wb = load_workbook(p, read_only=True, data_only=True)
        try:
            out = []
            for name in wb.sheetnames:
                ws = wb[name]
                out.append({"sheet_name": name, "max_row": ws.max_row, "max_col": ws.max_column})
            return _dump(out)
        finally:
            wb.close()

    @tool
    def preview_sheet(file_name: str, sheet_name: str, n_rows: Optional[int] = None) -> str:
        """Preview the first N raw rows of a sheet, keyed by column letter.

        Use this to classify an unknown sheet (is it the requirements sheet?
        a signal list? a communication matrix? something irrelevant?) and to
        find which row actually holds column headers -- headers are NOT
        guaranteed to be row 1 (title rows, merged banners, and blank rows
        before the real header are common). Rows are returned with their
        1-based row number and values keyed by column letter (A, B, C, ...)
        rather than by a guessed header name, since the header row is
        exactly what you're trying to determine here.

        Args:
            file_name: Exact file name, as returned by list_input_files (no
                directory -- this run's input directory is fixed).
            sheet_name: Exact sheet name (from list_workbook_sheets).
            n_rows: How many rows to preview from the top. Defaults to the
                configured SHEET_PREVIEW_ROWS.

        Returns:
            JSON {"sheet_name": str, "rows": [{"row": int, "cells": {"A": ..., "B": ...}}]}.
        """
        p = _resolve_within(root, file_name)
        if p is None:
            return _access_denied(file_name, root)
        n = n_rows or settings.SHEET_PREVIEW_ROWS
        wb = load_workbook(p, read_only=True, data_only=True)
        try:
            if sheet_name not in wb.sheetnames:
                return _dump({"error": f"Sheet '{sheet_name}' not found. Available: {wb.sheetnames}"})
            ws = wb[sheet_name]
            rows_out = []
            for row_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=min(n, ws.max_row or 0)), start=1):
                cells = {}
                for col_idx, cell in enumerate(row, start=1):
                    if cell.value is not None:
                        cells[get_column_letter(col_idx)] = _cell_str(cell.value)
                rows_out.append({"row": row_idx, "cells": cells})
            return _dump({"sheet_name": sheet_name, "rows": rows_out})
        finally:
            wb.close()

    @tool
    def read_sheet_range(
        file_name: str,
        sheet_name: str,
        start_row: int,
        end_row: int,
        columns: Optional[list[str]] = None,
    ) -> str:
        """Read a bounded range of rows from a sheet, keyed by column letter.

        Always read large sheets in bounded chunks with this tool rather than
        trying to read an entire 500-1000 row sheet at once -- the range is
        capped server-side (see MAX_ROWS_PER_READ) specifically to protect the
        context budget. Call preview_sheet first to know which column letters
        matter before narrowing with `columns`.

        Args:
            file_name: Exact file name, as returned by list_input_files (no
                directory -- this run's input directory is fixed).
            sheet_name: Exact sheet name.
            start_row: 1-based first row to read (inclusive).
            end_row: 1-based last row to read (inclusive). Will be clamped to
                start_row + MAX_ROWS_PER_READ - 1 if the requested range is
                too large -- check the returned "truncated" flag and page
                through with a follow-up call if so.
            columns: Optional list of column letters (e.g. ["A", "C", "F"])
                to restrict the returned cells to. Omit to get every
                populated column.

        Returns:
            JSON {"sheet_name": str, "start_row": int, "end_row": int,
            "truncated": bool, "rows": [{"row": int, "cells": {...}}]}.
        """
        p = _resolve_within(root, file_name)
        if p is None:
            return _access_denied(file_name, root)
        wb = load_workbook(p, read_only=True, data_only=True)
        try:
            if sheet_name not in wb.sheetnames:
                return _dump({"error": f"Sheet '{sheet_name}' not found. Available: {wb.sheetnames}"})
            ws = wb[sheet_name]
            max_row = ws.max_row or 0
            requested_end = min(end_row, max_row)
            capped_end = min(requested_end, start_row + settings.MAX_ROWS_PER_READ - 1)
            truncated = capped_end < requested_end
            col_filter = set(c.upper() for c in columns) if columns else None

            rows_out = []
            if start_row <= capped_end:
                # Indexed by `enumerate`/known column position rather than a
                # cell's own .row/.column: in openpyxl's read-only mode, a
                # blank cell comes back as an `EmptyCell` placeholder with
                # neither attribute, so deriving position from the cell
                # itself (as the previous version did) crashed on any row
                # with an empty cell and silently reported "row": null
                # whenever column A itself was blank.
                for row_idx, row in enumerate(ws.iter_rows(min_row=start_row, max_row=capped_end), start=start_row):
                    cells = {}
                    for col_idx, cell in enumerate(row, start=1):
                        letter = get_column_letter(col_idx)
                        if col_filter is not None and letter not in col_filter:
                            continue
                        if cell.value is not None:
                            cells[letter] = _cell_str(cell.value)
                    rows_out.append({"row": row_idx, "cells": cells})

            return _dump(
                {
                    "sheet_name": sheet_name,
                    "start_row": start_row,
                    "end_row": capped_end,
                    "truncated": truncated,
                    "rows": rows_out,
                }
            )
        finally:
            wb.close()

    @tool
    def search_sheet(
        file_name: str,
        sheet_name: str,
        query: str,
        columns: Optional[list[str]] = None,
        regex: bool = False,
        max_results: Optional[int] = None,
    ) -> str:
        """Search a sheet for rows containing a query string, without reading it in full.

        This is how you resolve a signal/command/parameter name against a
        supporting document: rather than reading an entire 500-row signal
        list, search it for the term(s) mentioned in the requirement text
        and inspect only the matching rows. Matching is case-insensitive
        substring by default; pass regex=True to use `query` as a regular
        expression instead.

        Args:
            file_name: Exact file name, as returned by list_input_files (no
                directory -- this run's input directory is fixed).
            sheet_name: Exact sheet name.
            query: Substring or (if regex=True) regular expression to search
                for.
            columns: Optional list of column letters to restrict the search
                to. Omit to search every populated cell in each row.
            regex: Treat `query` as a regular expression instead of a plain
                substring.
            max_results: Cap on returned matching rows. Defaults to the
                configured MAX_SEARCH_RESULTS.

        Returns:
            JSON {"sheet_name": str, "match_count": int, "truncated": bool,
            "rows": [{"row": int, "cells": {...}}]}.
        """
        p = _resolve_within(root, file_name)
        if p is None:
            return _access_denied(file_name, root)
        limit = max_results or settings.MAX_SEARCH_RESULTS
        wb = load_workbook(p, read_only=True, data_only=True)
        try:
            if sheet_name not in wb.sheetnames:
                return _dump({"error": f"Sheet '{sheet_name}' not found. Available: {wb.sheetnames}"})
            ws = wb[sheet_name]
            col_filter = set(c.upper() for c in columns) if columns else None

            if regex:
                try:
                    pattern = re.compile(query, re.IGNORECASE)
                except re.error as e:
                    return _dump({"error": f"Invalid regex: {e}"})
                matcher = lambda text: pattern.search(text) is not None
            else:
                needle = query.lower()
                matcher = lambda text: needle in text.lower()

            matches = []
            truncated = False
            # See read_sheet_range above: position comes from `enumerate`,
            # never from a cell's own .row/.column, since a blank cell in
            # openpyxl's read-only mode is an `EmptyCell` placeholder
            # lacking both.
            for row_idx, row in enumerate(ws.iter_rows(), start=1):
                cells = {}
                hit = False
                for col_idx, cell in enumerate(row, start=1):
                    letter = get_column_letter(col_idx)
                    if col_filter is not None and letter not in col_filter:
                        continue
                    if cell.value is not None:
                        cells[letter] = _cell_str(cell.value)
                        if matcher(str(cell.value)):
                            hit = True
                if hit:
                    if len(matches) >= limit:
                        truncated = True
                        break
                    matches.append({"row": row_idx, "cells": cells})

            return _dump(
                {
                    "sheet_name": sheet_name,
                    "match_count": len(matches),
                    "truncated": truncated,
                    "rows": matches,
                }
            )
        finally:
            wb.close()

    return [list_input_files, list_workbook_sheets, preview_sheet, read_sheet_range, search_sheet]


# ---------------------------------------------------------------------------
# Write tool (bound to one run's exact output file)
# ---------------------------------------------------------------------------


def build_write_tool(output_path: Path):
    """Build the write_output_workbook tool, bound to exactly one output file.

    There is no path parameter to get wrong or misuse: the destination is
    fixed at build time to `output_path` (this run's precomputed final
    workbook path), so the tool physically cannot write anywhere else. Call
    this once per run with that run's own output path.
    """
    out = Path(output_path).resolve()

    @tool
    def write_output_workbook(rows: list[dict]) -> str:
        """Write the final, QA-validated test cases to this run's output .xlsx file.

        This enforces the fixed SYS5 template column set (see
        config.settings.OUTPUT_COLUMNS) regardless of what keys the caller
        used internally -- pass a dict per row using the exact
        OUTPUT_COLUMNS names (case-insensitive, whitespace-tolerant matching
        is applied). Only call this once, after QA validation has passed,
        and only with the final row set. The destination file is fixed for
        this run -- there is no path to specify.

        The written sheet is also formatted for readability: styled/frozen
        header row, per-column widths sized to field type, wrapped body
        text, zebra striping, borders, and an autofilter -- cosmetic only,
        it never alters a cell's value.

        Args:
            rows: List of dicts, one per test case, keyed by column name
                (see config.settings.OUTPUT_COLUMNS for the exact expected
                names).

        Returns:
            On success: JSON {"output_path": str, "row_count": int,
            "warnings": [str, ...]}. `warnings` lists any row missing an
            expected column (written as blank) so the caller can decide
            whether that's acceptable.
            On failure: JSON {"error": str} -- the destination is left
            exactly as it was before this call (nothing partial or corrupt
            is ever written to it); report this as a real failure, never as
            success, and consider calling this tool again.
        """
        out.parent.mkdir(parents=True, exist_ok=True)

        wb = Workbook()
        ws = wb.active
        ws.title = settings.OUTPUT_SHEET_NAME
        ws.append(settings.OUTPUT_COLUMNS)

        warnings: list[str] = []
        for i, row in enumerate(rows):
            lookup = {str(k).strip().lower(): v for k, v in row.items()}
            line = []
            missing = []
            for col in settings.OUTPUT_COLUMNS:
                key = col.strip().lower()
                if key in lookup:
                    line.append(lookup[key])
                else:
                    line.append("")
                    missing.append(col)
            if missing:
                warnings.append(f"Row {i + 1}: missing columns {missing}, written blank")
            ws.append(line)

        _beautify_workbook(ws, settings.OUTPUT_COLUMNS, len(rows))

        # Save to a temp file in the same directory (so the final move below
        # is an atomic same-filesystem rename, never a partial copy), verify
        # it actually reopens as the workbook we just wrote, then swap it
        # into place. Writing `wb.save(out)` directly used to mean any
        # mid-save failure -- a locked/synced destination file (e.g. a
        # OneDrive-synced output folder on Windows), a disk hiccup, the file
        # still open in Excel -- could leave a truncated or empty file
        # sitting at `out`. That file still satisfies a plain `is_file()`
        # check, so the run got reported as successful while the "saved"
        # workbook was actually corrupt or missing its rows. Nothing below
        # ever touches `out` itself until the write is confirmed good.
        # Extension must stay `out`'s own (.xlsx) -- openpyxl's save/load
        # refuse an unrecognized suffix like a plain ".tmp".
        tmp_path = out.parent / f".{out.stem}.{uuid.uuid4().hex}{out.suffix}"
        try:
            wb.save(tmp_path)
            verify_wb = load_workbook(tmp_path, read_only=True)
            try:
                verify_ws = verify_wb[settings.OUTPUT_SHEET_NAME]
                if verify_ws.max_row != len(rows) + 1:
                    raise ValueError(
                        f"Saved workbook has {verify_ws.max_row} row(s), expected "
                        f"{len(rows) + 1} (header + {len(rows)} data rows)"
                    )
            finally:
                verify_wb.close()
            os.replace(tmp_path, out)
        except Exception as e:
            tmp_path.unlink(missing_ok=True)
            return _dump({"error": f"Failed to save output workbook to {out}: {e}"})

        return _dump({"output_path": str(out), "row_count": len(rows), "warnings": warnings})

    return write_output_workbook
